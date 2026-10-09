"""Fill the Code Flayer 2.0 template with CLINISCOPE content (python-pptx, editable output)."""
import copy
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.enum.text import PP_ALIGN

SRC, OUT, QR = sys.argv[1], sys.argv[2], sys.argv[3]

INDIGO, TEAL, NAVY, BODY, MUTED, WHITE, LAV = "5369DC", "2DA6A0", "162239", "394252", "767F90", "FFFFFF", "C0CAFF"
SITE = "https://anirudhleetcode-max.github.io/cliniscope-healthcare-contradiction-detector/"
REPO = "https://github.com/anirudhleetcode-max/cliniscope-healthcare-contradiction-detector"

prs = Presentation(SRC)
slides = list(prs.slides)


def sh(slide, name):
    for s in slide.shapes:
        if s.name == name:
            return s
    raise KeyError(name)


def clone(slide, src):
    el = copy.deepcopy(src._element)
    el.nvSpPr.cNvPr.set("id", str(max(int(x.shape_id) for x in slide.shapes) + 1))
    slide.shapes._spTree.append(el)
    return slide.shapes[-1]


def geom(s, x=None, y=None, w=None, h=None):
    if x is not None: s.left = Inches(x)
    if y is not None: s.top = Inches(y)
    if w is not None: s.width = Inches(w)
    if h is not None: s.height = Inches(h)


def fill(s, paras, size=None, color=None, bold=None, after=None, align=None, italic=None, bullet=False):
    """Replace a text box's paragraphs, cloning the template's first run style.

    paras: list of paragraphs; each is a str or a list of (text, overrides) runs.
    overrides keys: b, color, size, i, link
    """
    tf = s.text_frame
    txBody = tf._txBody
    p0 = txBody.find(qn("a:p"))
    r0 = p0.find(qn("a:r"))
    base_rpr = copy.deepcopy(r0.find(qn("a:rPr")))
    base_ppr = copy.deepcopy(p0.find(qn("a:pPr")))
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    for para in paras:
        runs = [(para, {})] if isinstance(para, str) else para
        p = txBody.makeelement(qn("a:p"), {})
        txBody.append(p)
        ppr = copy.deepcopy(base_ppr) if base_ppr is not None else p.makeelement(qn("a:pPr"), {})
        p.append(ppr)
        for text, o in runs:
            r = p.makeelement(qn("a:r"), {})
            p.append(r)
            r.append(copy.deepcopy(base_rpr))
            t = r.makeelement(qn("a:t"), {})
            t.text = text
            r.append(t)
        # use python-pptx objects for styling
    for para, pobj in zip(paras, tf.paragraphs):
        runs = [(para, {})] if isinstance(para, str) else para
        if after is not None:
            pobj.space_after = Pt(after)
        if align is not None:
            pobj.alignment = align
        for (text, o), robj in zip(runs, pobj.runs):
            f = robj.font
            sz = o.get("size", size)
            if sz: f.size = Pt(sz)
            b = o.get("b", bold)
            if b is not None: f.bold = b
            it = o.get("i", italic)
            if it is not None: f.italic = it
            c = o.get("color", color)
            if c: f.color.rgb = RGBColor.from_string(c)
            if o.get("link"):
                robj.hyperlink.address = o["link"]
        if bullet:
            ppr = pobj._p.get_or_add_pPr()
            ppr.set("marL", "228600"); ppr.set("indent", "-228600")
            for tag, attrs in (("a:buClr", None), ("a:buFont", {"typeface": "Arial"}), ("a:buChar", {"char": "\u2022"})):
                el = ppr.makeelement(qn(tag), attrs or {})
                if tag == "a:buClr":
                    el.append(el.makeelement(qn("a:srgbClr"), {"val": TEAL}))
                ppr.append(el)
    return s


def lbl(label, value, lc=INDIGO, vc=NAVY):
    return [(label + "  ", {"b": True, "color": lc}), (value, {"b": False, "color": vc})]


def bullets(items, labelc=NAVY):
    out = []
    for label, text in items:
        out.append([(label, {"b": True, "color": labelc}), (text, {"b": False})])
    return out


def card(slide, box, stripe, head, body, x, y, w, h, bx_pad=0.24):
    geom(sh(slide, box), x, y, w, h)
    if stripe:
        geom(sh(slide, stripe), x, y, None, h)
    geom(sh(slide, head), x + bx_pad, y + 0.18, w - 0.45, 0.35)
    geom(sh(slide, body), x + bx_pad, y + 0.6, w - 0.45, h - 0.75)


def footer(slide, n):
    for s in slide.shapes:
        if s.has_text_frame and s.text_frame.text.startswith("CODE FLAYER 2.0  •  EDITABLE"):
            fill(s, ["CODE FLAYER 2.0  •  CLINISCOPE  •  PS-11R3"])
        if s.has_text_frame and s.text_frame.text.strip() == f"0{n} / 08":
            geom(s, 11.75, 7.08, 0.9, 0.22)
            fill(s, [f"0{n} / 08"])
            s.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


# ---------------------------------------------------------------- slide 1
s = slides[0]
geom(sh(s, "TextBox 5"), w=8.3)
fill(sh(s, "TextBox 5"), ["CODE FLAYER 2.0  •  PROJECT PRESENTATION  •  PS-11R3"])
geom(sh(s, "TextBox 6"), 0.8, 1.45, 8.5, 1.45)
fill(sh(s, "TextBox 6"), [[("CLINISCOPE", {"size": 48})],
                          [("Healthcare Record Contradiction Detector", {"size": 22, "color": LAV})]])
geom(sh(s, "TextBox 7"), 0.84, 3.05, 8.5, 0.5)
fill(sh(s, "TextBox 7"), ["Find the conflict. Trace the evidence. Keep humans in control."], size=18, italic=True)
geom(sh(s, "Rectangle 8"), y=3.75)
geom(sh(s, "TextBox 9"), y=4.15)
fill(sh(s, "TextBox 9"), [[("TEAM  ", {}), ("[Add official team name]", {"b": False, "color": "A0AECC"})]])
geom(sh(s, "TextBox 10"), y=4.7)
fill(sh(s, "TextBox 10"), [[("Team leader: ", {}), ("[Add team leader]", {"color": "A0AECC"})],
                           [("Team members: ", {}), ("[Add member names]", {"color": "A0AECC"})]])
geom(sh(s, "TextBox 11"), w=8.4)
fill(sh(s, "TextBox 11"), ["Domain: Healthcare  |  Problem statement PS-11R3  |  Vishnu Institute of Technology, Bhimavaram"])

# ---------------------------------------------------------------- slide 2
s = slides[1]
footer(s, 2)
fill(sh(s, "TextBox 3"), ["PS-11R3 · Healthcare Record Contradiction Detector. The framing below is our own interpretation of the problem."])
H1, H2 = 1.95, 4.05
card(s, "Rounded Rectangle 7", "Rectangle 8", "TextBox 9", "TextBox 10", 0.7, 1.9, 5.85, 1.95)
card(s, "Rounded Rectangle 11", "Rectangle 12", "TextBox 13", "TextBox 14", 6.8, 1.9, 5.8, 1.95)
card(s, "Rounded Rectangle 15", "Rectangle 16", "TextBox 17", "TextBox 18", 0.7, 4.0, 5.85, 1.95)
card(s, "Rounded Rectangle 19", "Rectangle 20", "TextBox 21", "TextBox 22", 6.8, 4.0, 5.8, 1.95)
fill(sh(s, "TextBox 10"), ["Patient information is spread across discharge summaries, intake forms, "
                           "medication lists and lab reports. Allergies, doses, diagnoses and lab values "
                           "can differ between records, and finding those differences by hand is slow."], size=13)
fill(sh(s, "TextBox 13"), [[("02  TARGET USERS ", {}), ("(INTENDED)", {"b": False, "color": MUTED, "size": 12})]])
fill(sh(s, "TextBox 14"), bullets([("Clinicians", " reviewing patient records"),
                                   ("Medical-records / HIM teams", ""),
                                   ("Clinical quality & patient-safety", " reviewers"),
                                   ("Authorized staff", " reconciling records")]), size=13, after=2, bullet=True)
fill(sh(s, "TextBox 18"), bullets([("Fragmented: ", "related statements sit in separate documents"),
                                   ("Context-sensitive: ", "dates, negation and dose changes blur conflicts"),
                                   ("Hard to trace: ", "each discrepancy needs its original source")]), size=13, after=2, bullet=True)
fill(sh(s, "TextBox 22"), ["A warning alone is not enough. Reviewers need the original evidence, the context "
                           "(dates, history) to judge it, and a controlled way to record their final decision."], size=13)
geom(sh(s, "TextBox 23"), 0.85, 6.15, 11.7, 0.65)
fill(sh(s, "TextBox 23"), [lbl("ONE-LINE PROBLEM:", "How might we help healthcare reviewers identify and investigate "
                               "inconsistencies across medical documents, with clear source evidence and human oversight?")],
     size=14)

# ---------------------------------------------------------------- slide 3
s = slides[2]
footer(s, 3)
fill(sh(s, "TextBox 3"), ["An evidence-first review prototype: documents in, traceable findings out, decisions made by people."])
card(s, "Rounded Rectangle 7", "Rectangle 8", "TextBox 9", "TextBox 10", 0.7, 1.9, 7.25, 2.25)
card(s, "Rounded Rectangle 11", "Rectangle 12", "TextBox 13", "TextBox 14", 8.2, 1.9, 4.4, 2.25)
fill(sh(s, "TextBox 10"), ["CLINISCOPE extracts text from PDF, TXT and DOCX records, identifies allergy, medication, "
                           "diagnosis, lab, procedure and smoking statements with deterministic rules, compares them "
                           "across the documents of one case, and flags potential discrepancies with verified source "
                           "quotations for human review."], size=14)
fill(sh(s, "TextBox 14"), ["Every finding is traceable: it is created only if each quote is found again at its exact "
                           "source offsets. Dated or documented changes are separated from true conflicts, and the "
                           "final decision stays with the reviewer."], size=12.5)
FW, FX = 3.65, [0.7, 4.8, 8.9]
feats = [("Rounded Rectangle 15", "Rectangle 16", "TextBox 17", "TextBox 18"),
         ("Rounded Rectangle 19", "Rectangle 20", "TextBox 21", "TextBox 22"),
         ("Rounded Rectangle 23", "Rectangle 24", "TextBox 25", "TextBox 26")]
for (b, st, hd, bd), x in zip(feats, FX):
    card(s, b, st, hd, bd, x, 4.4, FW, 1.6)
fill(sh(s, "TextBox 17"), ["01  INGEST DOCUMENTS"])
fill(sh(s, "TextBox 21"), ["02  DETECT WITH EVIDENCE"])
fill(sh(s, "TextBox 25"), ["03  REVIEW & AUDIT"])
fill(sh(s, "TextBox 18"), ["PDF (text layer), TXT and DOCX are validated, de-duplicated and extracted in the browser."], size=12.5)
fill(sh(s, "TextBox 22"), ["Rules compare statements and label each difference by type, backed by quoted evidence."], size=12.5)
fill(sh(s, "TextBox 26"), ["Confirm, resolve or dismiss with a required reason; every action lands in an audit log."], size=12.5)
# flow arrows between the feature cards (documents -> findings -> human review)
arrow_src = sh(slides[3], "TextBox 10")
for x in (4.35, 8.45):
    a = clone(s, arrow_src)
    a.name = "Flow Arrow"
    geom(a, x + 0.02, 4.95, 0.45, 0.4)
geom(sh(s, "TextBox 27"), 0.85, 6.2, 11.8, 0.6)
fill(sh(s, "TextBox 27"), [lbl("EXPECTED BENEFIT:", "Potential inconsistencies become easier to locate and investigate, "
                               "while evidence traceability and human judgment are preserved.")], size=14)

# ---------------------------------------------------------------- slide 4
s = slides[3]
footer(s, 4)
fill(sh(s, "TextBox 3"), ["The whole pipeline runs in the reviewer's browser: there is no backend server, no AI/LLM and no OCR in the current build."])
stages = [("Rounded Rectangle 7", "TextBox 8", "TextBox 9", "USER INPUT", ["PDF (text layer)", "TXT · DOCX", "Synthetic records"]),
          ("Rounded Rectangle 11", "TextBox 12", "TextBox 13", "FRONTEND", ["React + TypeScript", "Library · Queue", "Finding detail"]),
          ("Rounded Rectangle 15", "TextBox 16", "TextBox 17", "CLIENT SERVICES", ["Validate · dedupe", "Analyze · review", "IndexedDB (Dexie)"]),
          ("Rounded Rectangle 19", "TextBox 20", "TextBox 21", "PROCESSING", ["Statement rules", "Conflict rules", "Evidence check"]),
          ("Rounded Rectangle 23", "TextBox 24", "TextBox 25", "OUTPUT", ["Classified findings", "Quoted evidence", "Review + audit log"])]
for i, (box, head, sub, title, lines) in enumerate(stages):
    x = 0.75 + i * 2.5
    geom(sh(s, box), x, 1.95, 1.8, 1.45)
    geom(sh(s, head), x + 0.07, 2.08, 1.66, 0.28)
    fill(sh(s, head), [title])
    geom(sh(s, sub), x + 0.07, 2.45, 1.66, 0.85)
    fill(sh(s, sub), lines, size=10.5, after=1)
for name, x in zip(["TextBox 10", "TextBox 14", "TextBox 18", "TextBox 22"], [2.63, 5.13, 7.63, 10.13]):
    geom(sh(s, name), x, 2.42)
card(s, "Rounded Rectangle 26", "Rectangle 27", "TextBox 28", "TextBox 29", 0.75, 3.7, 5.85, 2.3)
card(s, "Rounded Rectangle 30", "Rectangle 31", "TextBox 32", "TextBox 33", 6.8, 3.7, 5.8, 2.3)
fill(sh(s, "TextBox 29"), [[("1  Extraction: ", {"b": True, "color": NAVY}), ("pdf.js, mammoth, UTF-8 text with offsets", {})],
                           [("2  Statements: ", {"b": True, "color": NAVY}), ("negation, hedging, history, dose changes", {})],
                           [("3  Detection: ", {"b": True, "color": NAVY}), ("per-case comparison, units & dates", {})],
                           [("4  Evidence: ", {"b": True, "color": NAVY}), ("quotes re-verified; View in source", {})],
                           [("5  Review: ", {"b": True, "color": NAVY}), ("state machine, notes, audit log", {})]],
     size=12.5, after=2)
fill(sh(s, "TextBox 33"), bullets([("Uploads: ", "type, magic-byte and size checks; dedupe"),
                                   ("Findings: ", "only if every quote matches its source"),
                                   ("Dates: ", "specimen dates and dose changes reclassify"),
                                   ("Scanned PDFs: ", "flagged; every decision needs a reason")]),
     size=12.5, after=2, bullet=True)
geom(sh(s, "TextBox 34"), 0.85, 6.2, 11.7, 0.6)
fill(sh(s, "TextBox 34"), [lbl("ARCHITECTURE NOTE:", "Extraction, statement rules, detection and review are separate "
                               "TypeScript modules with unit tests, so each stage can be tested and replaced independently.")],
     size=14)

# ---------------------------------------------------------------- slide 5
s = slides[4]
footer(s, 5)
fill(sh(s, "TextBox 3"), ["Verified against package.json and the CI workflow. No Python, server framework, cloud database, OCR or LLM provider is used."])
cards5 = [("Rounded Rectangle 7", "Rectangle 8", "TextBox 9", "TextBox 10"),
          ("Rounded Rectangle 11", "Rectangle 12", "TextBox 13", "TextBox 14"),
          ("Rounded Rectangle 15", "Rectangle 16", "TextBox 17", "TextBox 18"),
          ("Rounded Rectangle 19", "Rectangle 20", "TextBox 21", "TextBox 22"),
          ("Rounded Rectangle 23", "Rectangle 24", "TextBox 25", "TextBox 26"),
          ("Rounded Rectangle 27", "Rectangle 28", "TextBox 29", "TextBox 30")]
content5 = [("FRONTEND", "React 18 + TypeScript, built with Vite and styled with Tailwind CSS."),
            ("BACKEND: NONE", "A TypeScript service layer in the browser runs upload, analysis and review."),
            ("STORAGE", "IndexedDB via Dexie: cases, documents, original files, findings and audit events."),
            ("CORE LOGIC (RULES)", "Deterministic TypeScript rules: vocabulary, negation, units and date handling."),
            ("DOCUMENT PROCESSING", "pdfjs-dist reads PDF text layers, mammoth reads DOCX; TXT read as UTF-8."),
            ("TESTING & DEPLOYMENT", "Vitest + Playwright in GitHub Actions; static build hosted on GitHub Pages.")]
for i, ((b, st, hd, bd), (t, d)) in enumerate(zip(cards5, content5)):
    x = [0.7, 4.75, 8.8][i % 3]
    y = 1.9 if i < 3 else 3.75
    card(s, b, st, hd, bd, x, y, 3.8, 1.6)
    fill(sh(s, hd), [t])
    fill(sh(s, bd), [d], size=13)
geom(sh(s, "TextBox 31"), 0.85, 5.65, 11.6, 0.9)
fill(sh(s, "TextBox 31"), [lbl("WHY THIS STACK?", "A lightweight, testable static web app: documents are processed and "
                               "stored in the reviewer's own browser, so no server or credentials are needed. A backend, OCR "
                               "or AI service would be added only when the implementation requires it.")], size=14)

# ---------------------------------------------------------------- slide 6
s = slides[5]
footer(s, 6)
fill(sh(s, "TextBox 3"), ["A working, tested prototype today, with a staged and cautious path toward real-world evaluation."])
for n in ("Rounded Rectangle 7", "Rounded Rectangle 8"):
    geom(sh(s, n), y=1.9, h=3.9)
for n in ("Rounded Rectangle 9", "Rounded Rectangle 10"):
    geom(sh(s, n), y=1.9)
for n in ("TextBox 11", "TextBox 12"):
    geom(sh(s, n), y=2.01)
geom(sh(s, "TextBox 13"), 0.95, 2.62, 5.4, 3.1)
geom(sh(s, "TextBox 14"), 7.05, 2.62, 5.35, 3.1)
fill(sh(s, "TextBox 13"), bullets([
    ("Tools & skills: ", "React, TypeScript, rule design, pdf.js / mammoth, Vitest and Playwright, all in use today."),
    ("Data & infrastructure: ", "synthetic demo records and browser storage only; no hospital systems or real patient data."),
    ("Status & constraints: ", "deployed prototype; 34 unit + 8 end-to-end tests pass. Limited vocabulary, no OCR, no login."),
    ("Risk → mitigation: ", "mis-extraction and context false positives → offset-verified quotes, uncertainty labels, human review, regression tests.")]),
    size=13.5, after=9, bullet=True)
fill(sh(s, "TextBox 14"), bullets([
    ("Users & value: ", "authorized reviewers who reconcile records and need evidence, not just alerts."),
    ("Adoption path: ", "synthetic prototype → representative test set → clinical reviewer feedback → security & privacy review → approved limited pilot."),
    ("Cost drivers: ", "hosting (static today), document processing, future storage or AI APIs, testing and maintenance."),
    ("Possible model (future): ", "institutional licensing or record-system integration via pilot partners. No partners or customers today.")]),
    size=13.5, after=9, bullet=True)
geom(sh(s, "Rounded Rectangle 15"), 0.95, 6.0, 11.4, 0.85)
geom(sh(s, "TextBox 16"), 1.15, 6.08, 11.0, 0.7)
fill(sh(s, "TextBox 16"), [lbl("KEY TAKEAWAY:", "CLINISCOPE builds on a working software prototype and offers a practical path toward "
                               "evaluated, evidence-based record review, subject to further validation and secure deployment.")],
     size=13)

# ---------------------------------------------------------------- slide 7
s = slides[6]
footer(s, 7)
fill(sh(s, "TextBox 3"), ["Scalability items and roadmap stages 02–04 are planned, not built; impact is a goal to be measured in evaluation."])
cards7 = [("Rounded Rectangle 7", "Rectangle 8", "TextBox 9", "TextBox 10"),
          ("Rounded Rectangle 11", "Rectangle 12", "TextBox 13", "TextBox 14"),
          ("Rounded Rectangle 15", "Rectangle 16", "TextBox 17", "TextBox 18")]
for (b, st, hd, bd), x in zip(cards7, [0.7, 4.75, 8.8]):
    card(s, b, st, hd, bd, x, 1.85, 3.8, 2.35)
fill(sh(s, "TextBox 13"), [[("SCALABILITY ", {}), ("(PLANNED)", {"b": False, "color": MUTED, "size": 12})]])
fill(sh(s, "TextBox 17"), [[("EXPECTED IMPACT ", {}), ("(GOALS)", {"b": False, "color": MUTED, "size": 12})]])
B = lambda items: [[(t, {})] for t in items]
fill(sh(s, "TextBox 10"), B(["Modular engine in src/lib", "Regression tests run in CI", "Versioned DB schema",
                             "Unsupported files flagged openly", "Synthetic, governed test data"]), size=12, after=1, bullet=True)
fill(sh(s, "TextBox 14"), B(["OCR for scanned records", "More formats and vocabulary", "Evidence-checked AI assistance",
                             "Secure backend, login, multi-user", "Record-system integration"]), size=12, after=1, bullet=True)
fill(sh(s, "TextBox 18"), B(["Aims to cut manual searching", "More traceable investigations", "Clearer view of source evidence",
                             "Structured review records", "Better awareness of uncertainty"]), size=12, after=1, bullet=True)
geom(sh(s, "TextBox 19"), y=4.35)
geom(sh(s, "Rectangle 20"), y=5.0)
road = [("Oval 21", "TextBox 22", "01  Prototype (now)", "Synthetic records: ingestion, evidence, review"),
        ("Oval 23", "TextBox 24", "02  Pilot testing", "Measure extraction, accuracy and usability"),
        ("Oval 25", "TextBox 26", "03  Improve", "OCR, context, evidence checks, security"),
        ("Oval 27", "TextBox 28", "04  Deploy & scale", "Only after accuracy, privacy and approvals")]
for i, (ov, tb, title, desc) in enumerate(road):
    x = 1.05 + i * 2.85
    geom(sh(s, ov), x, 4.81)
    geom(sh(s, tb), x - 0.05, 5.33, 2.65, 0.65)
    fill(sh(s, tb), [[(title, {})], [(desc, {"b": False, "size": 10.5, "color": BODY})]])
geom(sh(s, "Rectangle 20"), 1.0, 5.0, 9.8)
geom(sh(s, "TextBox 29"), 0.85, 6.1, 11.7, 0.75)
fill(sh(s, "TextBox 29"), [lbl("SUCCESS METRICS (to be measured):", "detection precision & recall on a labelled set  •  % of findings "
                               "with verifiable evidence  •  false-positive rate  •  reviewer agreement  •  median time per finding  "
                               "•  extraction success rate")], size=12)

# ---------------------------------------------------------------- slide 8
s = slides[7]
geom(sh(s, "TextBox 4"), 0.85, 2.25, 8.3, 1.3)
fill(sh(s, "TextBox 4"), ["CLINISCOPE helps reviewers identify potential inconsistencies across healthcare documents, "
                          "trace every finding to verified source evidence, and record decisions with human oversight."], size=18)
geom(sh(s, "TextBox 6"), w=8.3)
fill(sh(s, "TextBox 6"), ["Find the conflict. Trace the evidence. Keep humans in control."], size=18)
geom(sh(s, "TextBox 7"), y=4.95)
geom(sh(s, "TextBox 8"), 0.85, 5.55, 8.5, 1.2)
fill(sh(s, "TextBox 8"), [lbl("PS-11R3", "Healthcare Record Contradiction Detector", lc=LAV, vc="DEE3ED"),
                          [("Live demo:  ", {"b": True, "color": LAV}),
                           (SITE.replace("https://", "").rstrip("/"), {"color": "DEE3ED", "link": SITE})],
                          [("Code:  ", {"b": True, "color": LAV}),
                           (REPO.replace("https://", ""), {"color": "DEE3ED", "link": REPO})],
                          [("Contact:  ", {"b": True, "color": LAV}), ("[Add team email]", {"color": "A0AECC"})]],
     size=11, after=3)
fill(sh(s, "TextBox 11"), ["CLINISCOPE"])
geom(sh(s, "TextBox 11"), 9.9, 6.85, 3.2, 0.3)
pic = s.shapes.add_picture(QR, Inches(11.35), Inches(4.8), Inches(1.5), Inches(1.5))
pic.name = "QR Live Demo"
pic._element.nvPicPr.cNvPr.set("descr", "QR code linking to the CLINISCOPE live demo: " + SITE)
cap = clone(s, sh(s, "TextBox 11"))
cap.name = "QR Caption"
geom(cap, 9.9, 6.42, 3.2, 0.25)
fill(cap, ["SCAN FOR LIVE DEMO"], size=9, color=LAV)
# keep the two decorative squares, nudged up to make room for the QR code
geom(sh(s, "Rounded Rectangle 9"), 10.35, 1.0)
geom(sh(s, "Rounded Rectangle 10"), 11.15, 3.0)

# hyperlinks only appear on the dark closing slide: make the theme's link colours readable there
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
theme = prs.slide_master.part.part_related_by(RT.THEME)
blob = theme.blob.decode("utf-8")
blob = blob.replace('<a:hlink><a:srgbClr val="0000FF"/>', '<a:hlink><a:srgbClr val="C0CAFF"/>')
blob = __import__("re").sub(r'<a:folHlink><a:srgbClr val="[0-9A-Fa-f]{6}"/>', '<a:folHlink><a:srgbClr val="A0AECC"/>', blob)
theme._blob = blob.encode("utf-8")

prs.save(OUT)
print("saved", OUT)
